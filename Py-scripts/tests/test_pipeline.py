import json

import pandas as pd
import pytest
import requests

import sample_data as sd
from cve2at import cli, fetch
from cve2at.attack_tree_lib import TreeResult, validate_tree
from cve2at.config import Workspace, cpe_slug, validate_cpe
from cve2at.cve_parse import extract_cvss
from cve2at.errors import ConfigError, FetchError
from cve2at.tabulate import clean_weaknesses


def _tree_result(nodes, edges, leaf_count, root_id="root"):
    return TreeResult(
        name="test", dot=None, nodes_data=nodes, edges_data=edges,
        node_html={}, leaf_count=leaf_count, root_id=root_id,
    )


def test_validate_tree_accepts_a_well_formed_tree():
    nodes = [
        {"id": "root", "type": "gate", "label": "OR"},
        {"id": "and1", "type": "gate", "label": "AND"},
        {"id": "and2", "type": "gate", "label": "AND"},
        {"id": "tech1", "type": "tech", "label": "T1"},
        {"id": "capec1", "type": "capec", "label": "C1"},
        {"id": "tech2", "type": "tech", "label": "T2"},
        {"id": "capec2", "type": "capec", "label": "C2"},
    ]
    edges = [
        {"from": "root", "to": "and1"}, {"from": "root", "to": "and2"},
        {"from": "and1", "to": "tech1"}, {"from": "and1", "to": "capec1"},
        {"from": "and2", "to": "tech2"}, {"from": "and2", "to": "capec2"},
    ]
    result = validate_tree(_tree_result(nodes, edges, leaf_count=2))
    assert result["ok"] is True


def test_validate_tree_flags_a_malformed_and_gate():
    # AND gate with three children instead of exactly two.
    nodes = [
        {"id": "root", "type": "gate", "label": "OR"},
        {"id": "and1", "type": "gate", "label": "AND"},
        {"id": "and2", "type": "gate", "label": "AND"},
        {"id": "tech1", "type": "tech", "label": "T1"},
        {"id": "capec1", "type": "capec", "label": "C1"},
        {"id": "capec1b", "type": "capec", "label": "C1b"},
        {"id": "tech2", "type": "tech", "label": "T2"},
        {"id": "capec2", "type": "capec", "label": "C2"},
    ]
    edges = [
        {"from": "root", "to": "and1"}, {"from": "root", "to": "and2"},
        {"from": "and1", "to": "tech1"}, {"from": "and1", "to": "capec1"},
        {"from": "and1", "to": "capec1b"},
        {"from": "and2", "to": "tech2"}, {"from": "and2", "to": "capec2"},
    ]
    result = validate_tree(_tree_result(nodes, edges, leaf_count=2))
    assert result["ok"] is False
    assert result["and_gates_well_formed"] is False
    assert "and1" in result["malformed_and_gates"]


def test_validate_tree_flags_a_shared_multi_parent_node():
    nodes = [
        {"id": "root", "type": "gate", "label": "OR"},
        {"id": "and1", "type": "gate", "label": "AND"},
        {"id": "and2", "type": "gate", "label": "AND"},
        {"id": "tech1", "type": "tech", "label": "T1"},
        {"id": "capec1", "type": "capec", "label": "C1"},
        {"id": "tech2", "type": "tech", "label": "T2"},
    ]
    edges = [
        {"from": "root", "to": "and1"}, {"from": "root", "to": "and2"},
        {"from": "and1", "to": "tech1"}, {"from": "and1", "to": "capec1"},
        {"from": "and2", "to": "tech2"}, {"from": "and2", "to": "capec1"},  # shared node
    ]
    result = validate_tree(_tree_result(nodes, edges, leaf_count=2))
    assert result["ok"] is False
    assert result["tree_condition_ok"] is False
    assert "capec1" in result["nodes_with_multiple_parents"]


# small units
def test_validate_cpe_and_slug():
    assert cpe_slug(validate_cpe(sd.CPE)) == "acme_test_pump"
    assert cpe_slug("cpe:2.3:a:vendor:prod:1.2:*:*:*:*:*:*:*") == "vendor_prod_1.2"
    for bad in ["", "not-a-cpe", "cpe:2.3:h:only:three", "cpe:2.2:h:a:b:-:*:*:*:*:*:*:*",
                "cpe:2.3:x:a:b:-:*:*:*:*:*:*:*"]:
        with pytest.raises(ConfigError):
            validate_cpe(bad)


def test_extract_cvss_prefers_newest_version():
    v31 = {"cvssMetricV31": [{"cvssData": {"baseScore": 9.8}}], "cvssMetricV2": [{"cvssData": {"baseScore": 5}}]}
    assert extract_cvss(v31)["CVSS Version"] == "3.1"
    assert extract_cvss({"cvssMetricV2": [{"cvssData": {"baseScore": 5, "accessVector": "NETWORK"}}]})["attackVector"] == "NETWORK"
    assert extract_cvss({}) == {}


def test_clean_weaknesses():
    assert clean_weaknesses(["CWE-79", "NVD-CWE-Other", "CWE-79", "CWE-20"]) == "CWE-20; CWE-79"
    assert clean_weaknesses(["NVD-CWE-noinfo"]) == ""
    assert clean_weaknesses([]) == ""


# fetching
def test_fetch_pages_through_all_results(monkeypatch):
    calls = []
    monkeypatch.setattr(requests.Session, "get", sd.make_fake_get(sd.fake_vulnerabilities(), calls))
    vulns = fetch.fetch_cves_by_cpe(sd.CPE, results_per_page=3)
    assert len(vulns) == 8
    assert len(calls) == 3                                    # 3 + 3 + 2
    assert "cpeName=cpe%3A2.3%3Ah%3Aacme" in calls[0]["url"]  # CPE is URL-encoded


def test_fetch_sends_api_key(monkeypatch):
    calls = []
    monkeypatch.setattr(requests.Session, "get", sd.make_fake_get(sd.fake_vulnerabilities(), calls))
    fetch.fetch_cves_by_cpe(sd.CPE, api_key="secret")
    assert calls[0]["headers"] == {"apiKey": "secret"}


def test_fetch_retries_rate_limit_then_succeeds(monkeypatch):
    calls = []
    monkeypatch.setattr(requests.Session, "get",
                        sd.make_fake_get(sd.fake_vulnerabilities(), calls, fail_first=2, fail_status=403))
    assert len(fetch.fetch_cves_by_cpe(sd.CPE)) == 8
    assert len(calls) == 3


def test_fetch_gives_up_with_a_clear_error(monkeypatch):
    monkeypatch.setattr(requests.Session, "get",
                        sd.make_fake_get(sd.fake_vulnerabilities(), fail_first=99, fail_status=503))
    with pytest.raises(FetchError, match="503"):
        fetch.fetch_cves_by_cpe(sd.CPE, max_retries=2)


def test_fetch_splits_rejected_and_replaces_stale_files(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "ws").ensure()
    monkeypatch.setattr(requests.Session, "get", sd.make_fake_get(sd.fake_vulnerabilities()))
    ws.cve_dir.mkdir(); (ws.cve_dir / "CVE-1999-STALE.json").write_text("{}")   # leftover from an older run

    result = fetch.process_and_save_cves(sd.CPE, ws)
    assert result == {"total": 8, "accepted": 7, "rejected": 1}
    assert not (ws.cve_dir / "CVE-1999-STALE.json").exists()
    assert (ws.rejected_dir / "CVE-2021-0008.json").is_file()
    assert json.loads((ws.cve_dir / "CVE-2024-0001.json").read_text())["id"] == "CVE-2024-0001"


def test_failed_fetch_keeps_previous_data(tmp_path, monkeypatch):
    ws = Workspace(tmp_path / "ws").ensure()
    ws.cve_dir.mkdir(); (ws.cve_dir / "CVE-2000-0001.json").write_text("{}")
    monkeypatch.setattr(requests.Session, "get", sd.make_fake_get([], fail_first=99, fail_status=503))
    with pytest.raises(FetchError):
        fetch.process_and_save_cves(sd.CPE, ws, max_retries=0)
    assert (ws.cve_dir / "CVE-2000-0001.json").exists()


# whole pipeline
def _run(tmp_path, datasets, *extra):
    return cli.main(["run", sd.CPE, "-d", str(datasets), "-o", str(tmp_path / "out"), *extra])


def test_full_pipeline_end_to_end(tmp_path, datasets, fake_nvd, capsys):
    assert _run(tmp_path, datasets) == 0
    ws = tmp_path / "out" / "acme_test_pump"

    # every expected output file is there
    for f in ["cve_data.xlsx", "cve_data_formated.xlsx", "Device_Attack_Pattern.xlsx",
              "device_compromise_dataset.xlsx", "run_summary.json"]:
        assert (ws / f).is_file(), f
    assert len(list((ws / "cve_data").glob("*.json"))) == 7
    assert len(list((ws / "rejected_cve_data").glob("*.json"))) == 1
    assert (ws / "outputs/comprehensive_tree/comprehensive_attack_tree.html").is_file()
    assert (ws / "outputs/comprehensive_tree/comprehensive_attack_tree.png").is_file()
    for name in ["Initial_Access", "Execution", "Persistence", "Impact"]:
        for ext in ["svg", "html", "json", "dot"]:
            assert (ws / f"outputs/tactic_trees/{name}.{ext}").is_file(), (name, ext)

    # numbers known from the design of the synthetic data
    s = json.loads((ws / "run_summary.json").read_text())
    assert s["validation_ok"] is True
    assert s["stages"]["fetch"] == {"total": 8, "accepted": 7, "rejected": 1}
    assert s["stages"]["tactic_trees"]["total_paths"] == 9
    # default dedup mode is per_tactic, so the combined view equals the per-tactic total
    assert s["stages"]["comprehensive_tree"]["dedup_mode"] == "per_tactic"
    assert s["stages"]["comprehensive_tree"]["total_paths"] == 9
    assert s["stages"]["comprehensive_tree"]["per_tactic_total"] == 9
    assert s["stages"]["comprehensive_tree"]["earliest_tactic_total"] == 7
    assert s["stages"]["comprehensive_tree"]["cross_tactic_duplicates_removed"] == 0
    assert s["stages"]["comprehensive_tree"]["subtree_sum_matches"] is True
    assert s["stages"]["tactic_trees"]["cves_in_P"] is not None
    assert s["stages"]["tactic_trees"]["cwes_in_P"] >= 1

    # runtime and size figures
    assert s["scale"]["attack_matrix"] == {"tactics": 7, "techniques": 5, "sub_techniques": 2}
    assert s["scale"]["paths"] == 9
    assert s["scale"]["chains_in_P"] == s["stages"]["tactic_trees"]["rows_kept"]
    timing = s["timing_seconds"]
    assert list(timing) == ["fetch", "tabulate", "merge", "tactic_trees", "comprehensive_tree", "total"]
    assert all(v >= 0 for v in timing.values())
    assert timing["total"] == s["elapsed_seconds"]

    # padded ATT&CK sub-technique IDs survive the Excel round trip as text
    final = pd.read_excel(ws / "device_compromise_dataset.xlsx")
    assert "T1574.010" in set(final["Technique_ID_Only"])
    assert {"CWE-79", "CAPEC-63"} <= set(final["CWE ID"]) | set(final["CAPEC_ID"])
    # deprecated CAPEC-7 is excluded from the mapping
    assert final.loc[final["CAPEC_ID"] == "CAPEC-7", "capec_Name"].isna().all()

    assert "tactic" in capsys.readouterr().out.lower()


def test_dedup_mode_per_tactic_builds_the_bigger_tree(tmp_path, datasets, fake_nvd):
    assert _run(tmp_path, datasets, "--dedup-mode", "per_tactic", "--no-png") == 0
    ws = tmp_path / "out" / "acme_test_pump"
    s = json.loads((ws / "run_summary.json").read_text())
    assert s["stages"]["comprehensive_tree"]["total_paths"] == 9
    assert not (ws / "outputs/comprehensive_tree/comprehensive_attack_tree.png").exists()


def test_dedup_mode_earliest_tactic_is_kept_as_an_analysis_flag(tmp_path, datasets, fake_nvd):
    assert _run(tmp_path, datasets, "--dedup-mode", "earliest_tactic", "--no-png") == 0
    ws = tmp_path / "out" / "acme_test_pump"
    s = json.loads((ws / "run_summary.json").read_text())
    # per_tactic (9) minus the cross-tactic duplicates gives the earliest_tactic total (7)
    assert s["stages"]["comprehensive_tree"]["total_paths"] == 7
    assert s["stages"]["comprehensive_tree"]["cross_tactic_duplicates_removed"] == 2


def test_target_cve_restricts_the_analysis(tmp_path, datasets, fake_nvd):
    assert _run(tmp_path, datasets, "--no-png", "--target-cve", "CVE-2024-0001") == 0
    ws = tmp_path / "out" / "acme_test_pump"
    s = json.loads((ws / "run_summary.json").read_text())
    assert s["stages"]["merge"]["cves_in_scope"] == 1
    assert s["stages"]["tactic_trees"]["total_paths"] == 3        # CWE-79 -> 3 chains


def test_cpe_shorthand_and_skip_fetch(tmp_path, datasets, fake_nvd):
    assert cli.main([sd.CPE, "-d", str(datasets), "-o", str(tmp_path / "out"), "--no-png"]) == 0
    n_calls = len(fake_nvd)
    # second run reuses the downloaded files: no new NVD traffic
    assert cli.main([sd.CPE, "-d", str(datasets), "-o", str(tmp_path / "out"), "--no-png", "--skip-fetch"]) == 0
    assert len(fake_nvd) == n_calls
    summary = json.loads((tmp_path / "out" / "acme_test_pump" / "run_summary.json").read_text())
    assert summary["timing_seconds"]["fetch"] is None
    assert "fetch" not in summary["stages"]


def test_stage_commands_work_one_by_one(tmp_path, datasets, fake_nvd):
    base = ["-o", str(tmp_path / "out")]
    assert cli.main(["fetch", sd.CPE, *base]) == 0
    assert cli.main(["tabulate", sd.CPE, *base]) == 0
    assert cli.main(["merge", sd.CPE, *base, "-d", str(datasets)]) == 0
    assert cli.main(["trees", sd.CPE, *base, "--no-png", "--only", "comprehensive"]) == 0
    ws = tmp_path / "out" / "acme_test_pump"
    assert (ws / "outputs/comprehensive_tree/comprehensive_attack_tree.svg").is_file()
    assert not (ws / "outputs/tactic_trees").exists()


def atl_order(tactics):
    from cve2at.attack_tree_lib import mitre_order
    return mitre_order(tactics)


def test_validate_command_builds_the_workbook(tmp_path, datasets, fake_nvd):
    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"
    assert out.is_file()

    import openpyxl
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Consensus", "Results", "A1_annotation", "A2_annotation",
                             "Reference_Documents", "Labeling_Guide"]

    # both annotators get the same blank form
    for sheet in ("A1_annotation", "A2_annotation"):
        paths = wb[sheet]
        header = [c.value for c in paths[1]]
        assert header == ["Path ID", "Tactic", "CWE ID", "CWE Name", "Technique", "Technique Name",
                           "Sub-technique", "CAPEC ID", "CAPEC Name", "CVE(s)", "Path Key",
                           "Cross-Tactic Repeat",
                           "Validation", "Validation Rationale", "Notes"]
        assert paths.max_row - 1 == 9   # same 9 per-tactic paths as the tree stage
        for name in ("Validation", "Validation Rationale", "Notes"):
            col = header.index(name) + 1
            assert all(paths.cell(row=r, column=col).value in (None, "") for r in range(2, paths.max_row + 1))
    assert [c.value for c in wb["A1_annotation"][1]] == [c.value for c in wb["A2_annotation"][1]]
    paths = wb["A1_annotation"]
    header = [c.value for c in paths[1]]

    # there is no per-path document column: documents are cited in the rationale by Doc ID
    assert "Related Documents" not in header

    repeat_col = header.index("Cross-Tactic Repeat") + 1
    fills = {r: paths.cell(row=r, column=1).fill.fgColor.rgb for r in range(2, paths.max_row + 1)
             if paths.cell(row=r, column=repeat_col).value == "Yes"}
    assert len(fills) == 3
    assert len(set(fills.values())) == 1
    assert list(fills.values())[0] != "00000000"

    # the Consensus sheet takes both annotators' labels by formula and fills itself in where they agree
    cons = wb["Consensus"]
    assert [c.value for c in cons[1]] == ["Path ID", "Tactic", "A1 Validation", "A2 Validation",
                                          "A1 Rationale", "A2 Rationale", "Consensus Label",
                                          "Consensus Rationale", "Path Key", "Status"]
    assert cons.max_row - 1 == 9
    assert [cons.cell(row=r, column=1).value for r in range(2, 11)] == \
           [paths.cell(row=r, column=1).value for r in range(2, 11)]
    assert cons["C2"].value.startswith("=") and "A1_annotation" in cons["C2"].value
    assert cons["D2"].value.startswith("=") and "A2_annotation" in cons["D2"].value
    assert cons["G2"].value.startswith("=") and cons["J2"].value.startswith("=")

    # Results: the label table starts at A1, one row per tactic of the device, then Total and Percent
    res = wb["Results"]
    assert res["A1"].value == "Validation of the 9 paths (consensus labels)"
    assert [c.value for c in res[2]][:5] == ["Tactic", "Confirmed", "Plausible",
                                             "No supporting evidence found", "Total"]
    total_row = next(r for r in range(3, res.max_row + 1) if res.cell(row=r, column=1).value == "Total")
    tactics = [res.cell(row=r, column=1).value for r in range(3, total_row)]
    assert res.cell(row=total_row + 1, column=1).value == "Percent of all paths"
    assert "Impact" in tactics and all(tactics)
    assert tactics == atl_order(tactics)      # kill-chain order, one row per tactic
    assert len(tactics) == len(set(tactics))
    assert any(res.cell(row=r, column=1).value == "Cohen's kappa" for r in range(1, res.max_row + 1))

    docs = wb["Reference_Documents"]
    doc_header = [c.value for c in docs[1]]
    assert doc_header == ["Doc ID", "Advisory Tag", "Reference URL", "Alternate URLs", "Paths Citing"]
    assert docs.max_row > 1
    url_col = doc_header.index("Reference URL") + 1
    urls = [docs.cell(row=r, column=url_col).value for r in range(2, docs.max_row + 1)]
    assert all(urls)
    # Reference_Documents has one row per distinct Reference URL, however many paths or CVEs cite
    # it, never one row per (path, CVE, reference) triple
    assert len(urls) == len(set(urls))

    # Doc IDs are unique and sequential; Paths Citing is a formula keyed on the Doc ID
    doc_ids = [docs.cell(row=r, column=1).value for r in range(2, docs.max_row + 1)]
    assert len(doc_ids) == len(set(doc_ids))
    assert doc_ids == [f"D{i:04d}" for i in range(1, len(doc_ids) + 1)]
    cite_col = doc_header.index("Paths Citing") + 1
    assert all(str(docs.cell(row=r, column=cite_col).value).startswith("=")
               for r in range(2, docs.max_row + 1))

    guide = wb["Labeling_Guide"]
    assert [c.value for c in guide[1]] == ["Term", "Definition"]
    assert guide.max_row > 1


def test_validate_does_not_silently_reattach_judgments_when_a_path_id_is_reused(
        tmp_path, datasets, fake_nvd):
    """Path IDs come from enumeration order. If that order shifts between runs (a new CVE, a
    refreshed dataset), P00N can mean a different chain than when a person coded its judgment.
    The old judgment must not be attached to the new path; it is dropped and reported.
    """
    import openpyxl
    from cve2at import validate_report as vr
    from cve2at.config import Workspace

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    # Hand-code a judgment on annotator A1's sheet, as a human would in Excel.
    wb = openpyxl.load_workbook(out)
    paths = wb["A1_annotation"]
    header = [c.value for c in paths[1]]
    pidx = {name: i + 1 for i, name in enumerate(header)}
    target_row = 2
    target_path_id = paths.cell(row=target_row, column=pidx["Path ID"]).value
    original_key = paths.cell(row=target_row, column=pidx["Path Key"]).value
    original_tactic = paths.cell(row=target_row, column=pidx["Tactic"]).value
    paths.cell(row=target_row, column=pidx["Validation"], value="Confirmed")
    paths.cell(row=target_row, column=pidx["Validation Rationale"], value="Coded by a human.")
    wb.save(out)

    ws = Workspace(ws_dir)
    _, key_cols, P, ordered_tactics = vr._load_curated(ws.compromise_dataset_xlsx)
    cve_refs = vr._load_cve_references(ws)
    existing = vr._load_existing_workbook(out)
    assert existing["judgments"]["A1"][target_path_id]["Validation"] == "Confirmed"
    assert existing["identity"]["A1"][target_path_id] == (original_key, original_tactic)

    # Simulate the enumeration order shifting (e.g. an upstream dataset/library change) by
    # reversing tactic order, then rebuild against the same existing (unchanged) workbook data.
    reordered = list(reversed(ordered_tactics))
    path_df, doc_df = vr._build_tables(P, key_cols, reordered, cve_refs, existing)

    drifted = path_df.attrs["drifted"]
    assert drifted, "reversing tactic order should have reassigned at least one Path ID"
    drifted_ids = {d["path_id"] for d in drifted}
    assert target_path_id in drifted_ids
    assert {d["sheet"] for d in drifted if d["path_id"] == target_path_id} == {"A1"}

    # the stale Path ID must not carry the old "Confirmed" judgment onto the chain it now points at
    new_row = path_df[path_df["Path ID"] == target_path_id].iloc[0]
    assert new_row["Path Key"] != original_key or new_row["Tactic"] != original_tactic
    assert new_row["A1 Validation"] == ""
    assert new_row["A1 Validation Rationale"] == ""

    # The lost judgment is not silently discarded either, it's recoverable from the report.
    lost = next(d for d in drifted if d["path_id"] == target_path_id and d["sheet"] == "A1")
    assert lost["old_judgment"]["Validation"] == "Confirmed"
    assert lost["old_identity"] == (original_key, original_tactic)


def test_document_shared_across_cross_tactic_repeats_is_one_row(tmp_path, datasets, fake_nvd):
    """CWE-732 / CAPEC-17 / T1574.010 sits under 3 tactics (Persistence, Privilege Escalation,
    Defense Evasion) and every CVE in the fixture cites the same single URL. Reference_Documents is
    deduplicated by URL, so however many paths repeat, it is still one row.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    wb = openpyxl.load_workbook(out)
    paths = wb["A1_annotation"]
    pidx = {c.value: i + 1 for i, c in enumerate(paths[1])}
    repeated_rows = [r for r in range(2, paths.max_row + 1)
                      if paths.cell(row=r, column=pidx["Cross-Tactic Repeat"]).value == "Yes"]
    assert len(repeated_rows) == 3

    docs = wb["Reference_Documents"]
    didx = {c.value: i + 1 for i, c in enumerate(docs[1])}
    urls = [docs.cell(row=r, column=didx["Reference URL"]).value for r in range(2, docs.max_row + 1)]
    assert urls == ["https://example.org/a"]     # one row for the document, however many paths repeat


def test_renamed_document_url_keeps_its_doc_id_on_the_next_run(tmp_path, datasets, fake_nvd, capsys):
    """A publisher moves a page. The annotators put the new address in Reference URL, the old one in
    Alternate URLs, and keep the Doc ID; NVD still cites the old address. The re-run must not create
    a second document, and a rationale citing D0001 must stay valid.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    wb = openpyxl.load_workbook(out)
    docs = wb["Reference_Documents"]
    assert docs["A2"].value == "D0001" and docs["C2"].value == "https://example.org/a"
    docs["D2"] = "https://example.org/a"                     # the old address, still what NVD cites
    docs["C2"] = "https://docs.example.net/moved/a"
    wb["A1_annotation"].cell(row=2, column=14, value="See D0001.")
    wb.save(out)
    capsys.readouterr()

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    assert "not on the Reference_Documents sheet" not in capsys.readouterr().err
    wb2 = openpyxl.load_workbook(out)
    rows = [[c.value for c in row][:4] for row in wb2["Reference_Documents"].iter_rows(min_row=2)]
    assert rows == [["D0001", rows[0][1], "https://docs.example.net/moved/a", "https://example.org/a"]]
    assert wb2["A1_annotation"].cell(row=2, column=14).value == "See D0001."


def test_unrecorded_rename_makes_a_new_document_and_warns(tmp_path, datasets, fake_nvd, capsys):
    """If the new address is typed but the old one is not recorded, the CVE data's address matches
    nothing: the run must say so instead of silently attaching it, and must not reuse or drop D0001.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    wb = openpyxl.load_workbook(out)
    wb["Reference_Documents"]["C2"] = "https://docs.example.net/moved/a"
    wb.save(out)
    capsys.readouterr()

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    err = capsys.readouterr().err
    assert "new document D0002" in err and "Alternate URLs" in err
    ids = [r[0].value for r in openpyxl.load_workbook(out)["Reference_Documents"].iter_rows(min_row=2)]
    assert ids == ["D0001", "D0002"] or ids == ["D0002", "D0001"]


def test_document_ids_are_stable_across_a_rerun_with_no_changes(tmp_path, datasets, fake_nvd):
    """Re-running `cve2at validate` on an unchanged workspace must reuse every Doc ID, otherwise
    Doc ID references in a returned annotator copy would go stale.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    docs1 = openpyxl.load_workbook(out)["Reference_Documents"]
    before = {docs1.cell(row=r, column=1).value: docs1.cell(row=r, column=3).value
              for r in range(2, docs1.max_row + 1)}

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    docs2 = openpyxl.load_workbook(out)["Reference_Documents"]
    after = {docs2.cell(row=r, column=1).value: docs2.cell(row=r, column=3).value
             for r in range(2, docs2.max_row + 1)}
    assert after == before


def test_hand_added_document_row_gets_an_id_on_the_next_run(tmp_path, datasets, fake_nvd):
    """An annotator adds a row with the URL filled in and Doc ID blank. The next
    `cve2at validate` run must give it a permanent D#### instead of dropping it.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    wb = openpyxl.load_workbook(out)
    docs = wb["Reference_Documents"]
    new_row = docs.max_row + 1
    docs.cell(row=new_row, column=1, value=None)  # Doc ID left blank, as a human would
    docs.cell(row=new_row, column=2, value="Analogous device advisory")
    docs.cell(row=new_row, column=3, value="https://example.com/hand-added-advisory")
    wb.save(out)

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    docs2 = openpyxl.load_workbook(out)["Reference_Documents"]
    urls = {docs2.cell(row=r, column=3).value: docs2.cell(row=r, column=1).value
            for r in range(2, docs2.max_row + 1)}
    added_id = urls.get("https://example.com/hand-added-advisory")
    assert added_id and added_id.startswith("D")


def test_rerun_keeps_both_annotators_and_typed_consensus_but_rebuilds_the_automatic_cells(
        tmp_path, datasets, fake_nvd):
    """Both annotators' labels and a consensus label typed on a disputed row survive a re-run.
    The automatic Consensus formulas are written again, not frozen into values.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    wb = openpyxl.load_workbook(out)
    for sheet, label in (("A1_annotation", "Plausible"), ("A2_annotation", "No supporting evidence found")):
        ws = wb[sheet]
        header = [c.value for c in ws[1]]
        ws.cell(row=2, column=header.index("Validation") + 1, value=label)
        ws.cell(row=2, column=header.index("Validation Rationale") + 1, value=f"{sheet} reason")
    cons = wb["Consensus"]
    cons["G2"] = "Plausible"                       # typed after the two annotators talked it through
    cons["H2"] = "Both accepted the D0001 reading."
    wb.save(out)

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    wb2 = openpyxl.load_workbook(out)
    assert wb2["A1_annotation"].cell(row=2, column=13).value == "Plausible"          # M
    assert wb2["A2_annotation"].cell(row=2, column=13).value == "No supporting evidence found"
    assert wb2["A2_annotation"].cell(row=2, column=14).value == "A2_annotation reason"   # N
    cons2 = wb2["Consensus"]
    assert cons2["G2"].value == "Plausible" and cons2["H2"].value == "Both accepted the D0001 reading."
    # untouched rows still carry the automatic formulas
    assert cons2["G3"].value.startswith("=") and cons2["H3"].value.startswith("=")
    assert cons2["C2"].value.startswith("=") and cons2["J2"].value.startswith("=")


def test_old_single_sheet_workbook_is_not_migrated_but_keeps_its_documents(
        tmp_path, datasets, fake_nvd, capsys):
    """A workbook from the single-annotator version does not say who wrote the labels, so they are
    not attached to A1 or A2. The warning says so and the document list still carries over.
    """
    import openpyxl

    assert _run(tmp_path, datasets, "--no-png") == 0
    ws_dir = tmp_path / "out" / "acme_test_pump"
    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    out = ws_dir / "validation" / "validation_paths.xlsx"

    new = openpyxl.load_workbook(out)
    doc_rows = [[c.value for c in row] for row in new["Reference_Documents"].iter_rows()]
    old = openpyxl.Workbook()
    paths = old.active
    paths.title = "Validation Paths"
    paths.append(["Path ID", "Tactic", "Path Key", "Validation", "Validation Rationale", "Notes"])
    paths.append(["P001", "Initial Access", "x", "Confirmed", "old", ""])
    docs = old.create_sheet("Reference Documents")
    for row in doc_rows:
        docs.append(row)
    old.save(out)

    assert cli.main(["validate", sd.CPE, "-o", str(tmp_path / "out")]) == 0
    assert "old single 'Validation Paths' sheet" in capsys.readouterr().err
    wb = openpyxl.load_workbook(out)
    assert wb["A1_annotation"].cell(row=2, column=13).value in (None, "")
    ids_before = [r[0] for r in doc_rows[1:]]
    ids_after = [wb["Reference_Documents"].cell(row=r, column=1).value
                 for r in range(2, wb["Reference_Documents"].max_row + 1)]
    assert ids_after == ids_before


# friendly failures
def test_bad_cpe_is_a_clean_error(tmp_path, datasets, capsys):
    assert cli.main(["run", "nonsense", "-d", str(datasets), "-o", str(tmp_path)]) == 1
    assert "not a valid CPE" in capsys.readouterr().err


def test_missing_reference_datasets_fail_before_any_download(tmp_path, fake_nvd, capsys):
    empty = tmp_path / "empty"; empty.mkdir()
    assert cli.main(["run", sd.CPE, "-d", str(empty), "-o", str(tmp_path / "out")]) == 1
    assert "cwe_data.xlsx" in capsys.readouterr().err
    assert fake_nvd == []                                     # nothing was fetched


def test_cpe_with_no_cves(tmp_path, datasets, monkeypatch, capsys):
    monkeypatch.setattr(requests.Session, "get", sd.make_fake_get([]))
    assert cli.main(["run", sd.CPE, "-d", str(datasets), "-o", str(tmp_path / "out")]) == 1
    assert "no CVEs" in capsys.readouterr().err


def test_doctor(datasets, tmp_path, capsys):
    assert cli.main(["doctor", "-d", str(datasets)]) == 0
    assert cli.main(["doctor", "-d", str(tmp_path / "nope")]) == 1

def test_reproduce_fetches_and_runs_every_step(tmp_path, datasets, fake_nvd):
    out = tmp_path / "out"
    assert cli.main(["reproduce", sd.CPE, "-d", str(datasets), "-o", str(out), "--no-png"]) == 0
    assert fake_nvd, "reproduce must query NVD unless --skip-fetch is given"
    ws = next(out.iterdir())
    assert (ws / "cve_data").is_dir()
    assert (ws / "sensitivity.csv").is_file()
    assert (ws / "outputs" / "tables" / "cves.csv").is_file()
    assert (ws / "outputs" / "figures").is_dir()
    assert (ws / "outputs" / "insights" / "insights.json").is_file()


def test_reproduce_skip_fetch_reuses_files(tmp_path, datasets, fake_nvd):
    out = tmp_path / "out"
    assert cli.main(["reproduce", sd.CPE, "-d", str(datasets), "-o", str(out), "--no-png"]) == 0
    calls = len(fake_nvd)
    assert cli.main(["reproduce", sd.CPE, "-d", str(datasets), "-o", str(out), "--no-png", "--skip-fetch"]) == 0
    assert len(fake_nvd) == calls
