"""Stage 2b: the formatted review workbook (cve_data_formated.xlsx).

Sheets: CVE Summary, Affected Components, References, Analytics.
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import console
from .config import Workspace
from .cve_parse import cve_json_files, extract_cvss, english_description, load_cve, parse_cpe
from .errors import PipelineError

C_HEADER_BG, C_HEADER_FG, C_ROW_ALT, C_BORDER = "1A2B4A", "FFFFFF", "F0F4F8", "C8D0DC"
SEVERITY_COLORS = {"CRITICAL": "7B0000", "HIGH": "C0392B", "MEDIUM": "D4860A",
                   "LOW": "2E7D32", "NONE": "558B2F"}

_thin = Side(border_style="thin", color=C_BORDER)
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

MAIN_COLS = ["CVE ID", "Published", "Last Modified", "Status", "Source", "Description",
             "Primary Weakness (CWE)", "Secondary Weakness (CWE)", "CVSS Version", "Base Score",
             "Severity", "Attack Vector", "Attack Complexity", "Privileges Required",
             "User Interaction", "Scope", "Confidentiality Impact", "Integrity Impact",
             "Availability Impact", "Exploitability Score", "Impact Score", "Vector String"]
CONFIG_COLS = ["CVE ID", "Vulnerable", "Vendor", "Product / Component", "Version", "CPE Criteria"]
REF_COLS = ["CVE ID", "URL", "Tags"]


def _style_header(cell):
    cell.font = Font(bold=True, color=C_HEADER_FG, name="Calibri", size=10)
    cell.fill = PatternFill("solid", fgColor=C_HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border


def _style_data(cell, alt=False):
    cell.font = Font(name="Calibri", size=10)
    if alt:
        cell.fill = PatternFill("solid", fgColor=C_ROW_ALT)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = _border


def _autofit(ws, min_w=10, max_w=50):
    for col in ws.columns:
        length = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, min_w), max_w)


def _write_sheet(ws, df, col_widths=None):
    ws.freeze_panes = "A2"
    for ci, h in enumerate(df.columns, 1):
        _style_header(ws.cell(row=1, column=ci, value=h))
    for ri, row in enumerate(df.itertuples(index=False), 2):
        for ci, val in enumerate(row, 1):
            _style_data(ws.cell(row=ri, column=ci, value=val), alt=(ri % 2 == 0))
    ws.auto_filter.ref = ws.dimensions      # after the data is in, so the filter spans the table
    ws.row_dimensions[1].height = 28
    if col_widths:
        for col, w in col_widths.items():
            ws.column_dimensions[col].width = w
    else:
        _autofit(ws)


def _collect(ws_dir):
    main_rows, config_rows, ref_rows = [], [], []
    for path in cve_json_files(ws_dir):
        data = load_cve(path)
        cve_id = data["id"]

        weak_primary = weak_secondary = "N/A"
        for w in data.get("weaknesses", []):
            val = next((d["value"] for d in w.get("description", []) if d.get("lang") == "en"), None)
            if val:
                if w.get("type") == "Primary":
                    weak_primary = val
                else:
                    weak_secondary = val

        cvss = extract_cvss(data.get("metrics", {}))
        main_rows.append({
            "CVE ID": cve_id,
            "Published": data.get("published", "")[:10],
            "Last Modified": data.get("lastModified", "")[:10],
            "Status": data.get("vulnStatus", "N/A"),
            "Source": data.get("sourceIdentifier", "N/A"),
            "Description": english_description(data),
            "Primary Weakness (CWE)": weak_primary,
            "Secondary Weakness (CWE)": weak_secondary,
            "CVSS Version": cvss.get("CVSS Version"),
            "Base Score": cvss.get("baseScore"),
            "Severity": cvss.get("baseSeverity"),
            "Attack Vector": cvss.get("attackVector"),
            "Attack Complexity": cvss.get("attackComplexity"),
            "Privileges Required": cvss.get("privilegesRequired"),
            "User Interaction": cvss.get("userInteraction"),
            "Scope": cvss.get("scope"),
            "Confidentiality Impact": cvss.get("confidentialityImpact"),
            "Integrity Impact": cvss.get("integrityImpact"),
            "Availability Impact": cvss.get("availabilityImpact"),
            "Exploitability Score": cvss.get("exploitabilityScore"),
            "Impact Score": cvss.get("impactScore"),
            "Vector String": cvss.get("vectorString"),
        })

        for cfg in data.get("configurations", []):
            for node in cfg.get("nodes", []):
                for cpe in node.get("cpeMatch", []):
                    vendor, product, version = parse_cpe(cpe.get("criteria", ""))
                    config_rows.append({
                        "CVE ID": cve_id, "Vulnerable": cpe.get("vulnerable", False),
                        "Vendor": vendor, "Product / Component": product, "Version": version,
                        "CPE Criteria": cpe.get("criteria", "N/A"),
                    })

        seen_urls = set()
        for ref in data.get("references", []):
            url = ref.get("url", "")
            if url in seen_urls:
                continue
            seen_urls.add(url)
            ref_rows.append({"CVE ID": cve_id, "URL": url, "Tags": ", ".join(ref.get("tags", []))})

    # explicit columns: a CVE set with no configurations/references must not break the sheets
    return (pd.DataFrame(main_rows, columns=MAIN_COLS),
            pd.DataFrame(config_rows, columns=CONFIG_COLS),
            pd.DataFrame(ref_rows, columns=REF_COLS))


def _h1(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = Font(bold=True, color=C_HEADER_FG, name="Calibri", size=12)
    c.fill = PatternFill("solid", fgColor=C_HEADER_BG)
    c.alignment = Alignment(horizontal="left", vertical="center")
    c.border = _border


def _kv(ws, row, lbl, val, col=1):
    lc = ws.cell(row=row, column=col, value=lbl)
    lc.font = Font(bold=True, name="Calibri", size=10)
    lc.fill = PatternFill("solid", fgColor="E8EDF3")
    lc.border = _border
    lc.alignment = Alignment(horizontal="left", vertical="center")
    vc = ws.cell(row=row, column=col + 1, value=val)
    vc.font = Font(name="Calibri", size=10)
    vc.border = _border
    vc.alignment = Alignment(horizontal="center", vertical="center")


def build_formatted_workbook(ws: Workspace) -> dict:
    df_main, df_config, df_ref = _collect(ws.cve_dir)
    if df_main.empty:
        raise PipelineError(f"No CVE json files in '{ws.cve_dir}'. Run the fetch stage first.")

    wb = Workbook()

    # Sheet 1: CVE Summary
    ws1 = wb.active
    ws1.title = "CVE Summary"
    _write_sheet(ws1, df_main, col_widths={
        "A": 18, "B": 12, "C": 14, "D": 14, "E": 28, "F": 60, "G": 20, "H": 22, "I": 10, "J": 10,
        "K": 10, "L": 15, "M": 16, "N": 18, "O": 16, "P": 10, "Q": 18, "R": 16, "S": 18, "T": 16,
        "U": 16, "V": 35})
    sev_col_idx = df_main.columns.get_loc("Severity") + 1
    for ri in range(2, len(df_main) + 2):
        cell = ws1.cell(row=ri, column=sev_col_idx)
        sev = str(cell.value or "").upper()
        if sev in SEVERITY_COLORS:
            cell.fill = PatternFill("solid", fgColor=SEVERITY_COLORS[sev])
            cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
    score_col = get_column_letter(df_main.columns.get_loc("Base Score") + 1)
    ws1.conditional_formatting.add(
        f"{score_col}2:{score_col}{len(df_main) + 1}",
        ColorScaleRule(start_type="num", start_value=0, start_color="63BE7B",
                       mid_type="num", mid_value=5, mid_color="FFEB84",
                       end_type="num", end_value=10, end_color="F8696B"))

    # Sheet 2: Affected Components
    ws2 = wb.create_sheet("Affected Components")
    _write_sheet(ws2, df_config, col_widths={"A": 18, "B": 12, "C": 20, "D": 30, "E": 12, "F": 55})
    for ri in range(2, len(df_config) + 2):
        cell = ws2.cell(row=ri, column=2)
        if cell.value is True:
            cell.value = "YES"
            cell.fill = PatternFill("solid", fgColor="FDECEA")
            cell.font = Font(bold=True, color="C0392B", name="Calibri", size=10)
        else:
            cell.value = "no"

    # Sheet 3: References
    ws3 = wb.create_sheet("References")
    _write_sheet(ws3, df_ref, col_widths={"A": 18, "B": 65, "C": 40})

    # Sheet 4: Analytics
    ws4 = wb.create_sheet("Analytics")
    ws4.sheet_view.showGridLines = False
    _h1(ws4, 1, 1, "\U0001F4CA  CVE Dataset \u2014 Analytics Overview")
    ws4.merge_cells("A1:F1")
    ws4.row_dimensions[1].height = 30

    vulnerable = df_config[df_config["Vulnerable"] == True]  # noqa: E712 (bool column)
    _kv(ws4, 3, "Total CVEs", len(df_main))
    _kv(ws4, 4, "Total Affected Components", len(vulnerable))
    _kv(ws4, 5, "Total References", len(df_ref))
    _kv(ws4, 6, "Date Range (Published)",
        f"{df_main['Published'].min()}  \u2192  {df_main['Published'].max()}")

    _h1(ws4, 9, 1, "Severity Breakdown")
    ws4.merge_cells("A9:B9")
    for i, (sev, cnt) in enumerate(df_main["Severity"].value_counts().items(), 10):
        c_lbl, c_cnt = ws4.cell(row=i, column=1, value=sev), ws4.cell(row=i, column=2, value=int(cnt))
        bg = SEVERITY_COLORS.get(str(sev).upper(), "888888")
        for c in (c_lbl, c_cnt):
            c.fill = PatternFill("solid", fgColor=bg)
            c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
            c.border = _border
            c.alignment = Alignment(horizontal="center", vertical="center")

    _h1(ws4, 9, 4, "Attack Vector Breakdown")
    ws4.merge_cells("D9:E9")
    for i, (av, cnt) in enumerate(df_main["Attack Vector"].value_counts().items(), 10):
        c_lbl, c_cnt = ws4.cell(row=i, column=4, value=av), ws4.cell(row=i, column=5, value=int(cnt))
        for c in (c_lbl, c_cnt):
            c.fill = PatternFill("solid", fgColor="2C3E50")
            c.font = Font(color="FFFFFF", name="Calibri", size=10, bold=True)
            c.border = _border
            c.alignment = Alignment(horizontal="center", vertical="center")

    _h1(ws4, 20, 1, "Top Affected Vendors (vulnerable components)")
    ws4.merge_cells("A20:C20")
    for i, (v, cnt) in enumerate(vulnerable["Vendor"].value_counts().head(10).items(), 21):
        c_lbl, c_cnt = ws4.cell(row=i, column=1, value=v), ws4.cell(row=i, column=2, value=int(cnt))
        c_lbl.font = Font(name="Calibri", size=10)
        c_cnt.font = Font(name="Calibri", size=10, bold=True)
        c_cnt.fill = PatternFill("solid", fgColor="DDEEFF")
        for c in (c_lbl, c_cnt):
            c.border = _border
            c.alignment = Alignment(horizontal="left", vertical="center")

    _h1(ws4, 20, 4, "Score Statistics")
    ws4.merge_cells("D20:E20")
    scores = df_main["Base Score"].dropna()
    stats = ([("Average Base Score", round(float(scores.mean()), 2)), ("Max Base Score", float(scores.max())),
              ("Min Base Score", float(scores.min())), ("Median Base Score", float(scores.median()))]
             if len(scores) else [(k, "N/A") for k in
                                  ("Average Base Score", "Max Base Score", "Min Base Score", "Median Base Score")])
    for i, (lbl, val) in enumerate(stats, 21):
        lc = ws4.cell(row=i, column=4, value=lbl)
        lc.border = _border
        lc.font = Font(name="Calibri", size=10)
        lc.alignment = Alignment(horizontal="left")
        c = ws4.cell(row=i, column=5, value=val)
        c.border = _border
        c.font = Font(bold=True, name="Calibri", size=10)
        c.alignment = Alignment(horizontal="center")

    for col, w in [("A", 28), ("B", 12), ("C", 12), ("D", 32), ("E", 12)]:
        ws4.column_dimensions[col].width = w

    wb.save(ws.formatted_xlsx)
    console.info(f"Saved: {ws.formatted_xlsx}")
    console.info(f"Main rows: {len(df_main)}  |  Config rows: {len(df_config)}  |  Ref rows: {len(df_ref)}")
    return {"cves": len(df_main), "components": len(df_config), "references": len(df_ref)}
